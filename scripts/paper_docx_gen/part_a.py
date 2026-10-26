from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5); sec.page_height = Inches(11)
sec.left_margin = Inches(0.75); sec.right_margin = Inches(0.75)
sec.top_margin = Inches(0.75); sec.bottom_margin = Inches(0.75)
sty = doc.styles["Normal"]
sty.font.name = "Times New Roman"; sty.font.size = Pt(10)

def H1(t):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(4)
    r = p.add_run(t.upper()); r.bold=True; r.font.size=Pt(10); r.font.name="Times New Roman"
def H2(t):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(2)
    r = p.add_run(t); r.italic=True; r.font.size=Pt(10); r.font.name="Times New Roman"
def BP(text):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(3)
    p.paragraph_format.first_line_indent=Pt(14)
    r = p.add_run(text); r.font.size=Pt(10); r.font.name="Times New Roman"
def TC(t):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(6); p.paragraph_format.space_after=Pt(2)
    r = p.add_run(t); r.bold=True; r.font.size=Pt(8); r.font.name="Times New Roman"
def FT(tb):
    for row in tb.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                para.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for run in para.runs:
                    run.font.size=Pt(7.5); run.font.name="Times New Roman"
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
def TH(tb, hdrs, fs=8):
    for i,h in enumerate(hdrs):
        c=tb.rows[0].cells[i]; c.text=h
        for run in c.paragraphs[0].runs:
            run.bold=True; run.font.size=Pt(fs); run.font.name="Times New Roman"
def TR(tb, rows, fs=7.5):
    for rd in rows:
        row=tb.add_row()
        for i,v in enumerate(rd):
            c=row.cells[i]; c.text=v
            for run in c.paragraphs[0].runs:
                run.font.size=Pt(fs); run.font.name="Times New Roman"

# ---- TITLE ----
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(8)
r=p.add_run("Ssense: A Domain-Adapted Legal Small Language Model and Scalable Inference Architecture for Auditing Global Web Compliance under India's DPDP Act, 2023 and DPDP Rules, 2025")
r.bold=True; r.font.size=Pt(16); r.font.name="Times New Roman"
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(2)
r=p.add_run("Anubhav Das and Priyanshu Roy"); r.bold=True; r.font.size=Pt(11); r.font.name="Times New Roman"
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(8)
r=p.add_run("Department of Computer Science and Engineering, Institute of Engineering and Management, Kolkata, India")
r.italic=True; r.font.size=Pt(9); r.font.name="Times New Roman"
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(6)
r=p.add_run("Emails: anubhav.das@iemcal.com; priyanshu.roy@iemcal.com")
r.font.size=Pt(9); r.font.name="Times New Roman"

# ---- ABSTRACT ----
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY; p.paragraph_format.space_after=Pt(4)
r1=p.add_run("Abstract"); r1.bold=True; r1.italic=True; r1.font.size=Pt(10); r1.font.name="Times New Roman"
r2=p.add_run(u"\u2014Sovereign data protection legislation increasingly asserts extraterritorial jurisdiction over the global digital economy. Under Section 3(b) of India's Digital Personal Data Protection (DPDP) Act, 2023, data fiduciaries operating anywhere in the world must comply with statutory mandates when processing personal data of Indian data principals. In practice, global web platforms routinely deploy generic, boilerplate, or deceptive privacy policies that mask statutory non-compliance, while citizens and regulatory monitors lack automated tools to evaluate legal contracts at scale. General-purpose Large Language Models (LLMs) suffer from statutory hallucinations (58-88% error rates), cross-jurisdictional contamination (superimposing GDPR doctrines onto Indian law), and an epistemic failure to distinguish statutory silence from active breach. We present Ssense, the first domain-adapted Legal Small Language Model (Law-LM) and scalable inference architecture specifically fine-tuned and preference-aligned on the DPDP Act, 2023 and DPDP Rules, 2025. Built upon Qwen2.5-7B-Instruct, Ssense establishes a dual-track parameter-efficient architecture: a forensic statutory auditor (rank 128) and a citizen-facing legal dialogue model (rank 64), trained via rank-stabilized LoRA (rsLoRA) and Simple Preference Optimization (SimPO) over statutory reasoning traces, released as Ssense-DPDP-7B. We computationally model the DPDP Act by pruning 44 legislative sections into a 26-node actionable statutory decision tree. We introduce a three-step deductive reasoning chain and an omission firewall that discards findings derived from unstated practices, and couple the SLM with a deterministic Statutory Compliance Layer (SCL) as a neuro-symbolic backstop (V_final = V_SLM union V_SCL). Our inference architecture integrates PagedAttention continuous batching, prompt KV cache pre-warming, speculative decoding (1.8x speedup), five-tier SHA-256 cache hierarchy, and XGrammar pushdown CFG decoding (100% schema validity). Empirical audits of ten real-world Indian fiduciaries (>500M MAU) yielded zero statutory hallucinations, zero false positives, and 18/18 certification axes cleared at 95% Wilson confidence bounds.")
r2.font.size=Pt(10); r2.font.name="Times New Roman"
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY; p.paragraph_format.space_after=Pt(6)
r1=p.add_run("Index Terms"); r1.bold=True; r1.italic=True; r1.font.size=Pt(10); r1.font.name="Times New Roman"
r2=p.add_run(u"\u2014Legal Language Models, Statutory Alignment, Open-Weights Models, DPDP Act 2023, DPDP Rules 2025, Global Web Compliance, Scalable Inference Serving, Deontic Logic, Neuro-Symbolic AI, Privacy Policy Analysis")
r2.font.size=Pt(10); r2.font.name="Times New Roman"

print("Part A OK")
