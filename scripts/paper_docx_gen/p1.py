import os
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5); sec.page_height = Inches(11)
sec.left_margin = Inches(0.75); sec.right_margin = Inches(0.75)
sec.top_margin = Inches(0.75); sec.bottom_margin = Inches(0.75)
sty = doc.styles["Normal"]
sty.font.name = "Times New Roman"; sty.font.size = Pt(10)
def H1(t):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(4)
    r = p.add_run(t.upper()); r.bold=True; r.font.size=Pt(10); r.font.name="Times New Roman"
def H2(t):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(2)
    r = p.add_run(t); r.italic=True; r.font.size=Pt(10); r.font.name="Times New Roman"
def BP(text):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(3)
    p.paragraph_format.first_line_indent=Pt(14)
    r = p.add_run(text); r.font.size=Pt(10); r.font.name="Times New Roman"
def TC(t):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(6); p.paragraph_format.space_after=Pt(2)
    r = p.add_run(t); r.bold=True; r.font.size=Pt(8); r.font.name="Times New Roman"
def FT(tb):
    for row in tb.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                para.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for run in para.runs:
                    run.font.size=Pt(7.5); run.font.name="Times New Roman"
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
def TH(tb, hdrs, fs=8):
    for i,h in enumerate(hdrs):
        if i < len(tb.rows[0].cells):
            c=tb.rows[0].cells[i]; c.text=h
            for run in c.paragraphs[0].runs:
                run.bold=True; run.font.size=Pt(fs); run.font.name="Times New Roman"
def TR(tb, data, fs=7.5):
    for rd in data:
        row=tb.add_row()
        for i,v in enumerate(rd):
            if i < len(row.cells):
                c=row.cells[i]; c.text=str(v)
                for run in c.paragraphs[0].runs:
                    run.font.size=Pt(fs); run.font.name="Times New Roman"
print("helpers ok")
